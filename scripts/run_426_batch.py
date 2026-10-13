# -*- coding: utf-8 -*-
"""季节426 批量新建：读 9月确认消费券黑名id.xlsx（A=ID B=款号 C=季节）中季节=426 的行，
复用 tmall-clone-id-publisher 主流程 process_row，新ID写回 D 列、状态写 E 列。
用法: python run_426_batch.py [--start 229] [--end 233] [--yes] [--report x.json]
"""
import argparse
import json
import sys
from datetime import datetime
from pathlib import Path

sys.path.insert(0, r'E:\巴拉知识库\.agents\skills\tmall-clone-id-publisher\scripts')
from clone_id_flow import process_row, open_onsale_tab, build_pre_switch_set  # noqa: E402
import openpyxl  # noqa: E402

DEFAULT_SRC = r'E:\巴拉巴拉\AI\ID克隆项目\9月确认消费券黑名id.xlsx'  # 本机默认；他人使用请传 --src


def load_rows(path: str, start_row=None, end_row=None):
    wb = openpyxl.load_workbook(path)
    ws = wb['Sheet1']
    rows = []
    for r in range(2, ws.max_row + 1):
        if start_row and r < start_row:
            continue
        if end_row and r > end_row:
            continue
        season = ws.cell(r, 3).value
        oid = ws.cell(r, 1).value
        hh = ws.cell(r, 2).value
        if oid is None:
            continue
        if season is None or '426' not in str(season):
            continue
        rows.append({'excel_row': r, 'origin_id': str(oid).strip(),
                     'huohao': str(hh).strip() if hh else ''})
    return rows, wb, ws


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--src', default=DEFAULT_SRC, help='源Excel路径（默认本机路径）')
    ap.add_argument('--start', type=int, default=None)
    ap.add_argument('--end', type=int, default=None)
    ap.add_argument('--yes', action='store_true', help='授权提交')
    ap.add_argument('--cdp-url', default='http://127.0.0.1:9223')
    ap.add_argument('--report', default='')
    args = ap.parse_args()

    rows, wb, ws = load_rows(args.src, args.start, args.end)
    print(f'[load] {len(rows)} rows (season=426)')
    if not rows:
        print('no rows, exit')
        return

    pre_switch = build_pre_switch_set(rows)
    print(f'[pre-switch] 频次>=5 预切货号: {sorted(pre_switch) or "无"}')

    ctx = {'brand': '巴拉巴拉', 'seq_map': {}, 'pre_switch_set': pre_switch,
           'learned_set': set(), 'dry_run': False, 'yes': args.yes,
           'cdp_url': args.cdp_url}

    tab_onsale = open_onsale_tab(args.cdp_url)
    print('[tab]', tab_onsale.url[:80])

    results = []
    for row in rows:
        print(f"\n=== {row['origin_id']} (货号 {row['huohao']}, Excel行{row['excel_row']}) ===")
        try:
            r = process_row(tab_onsale, row, ctx)
        except Exception as e:
            r = {'excel_row': row['excel_row'], 'origin_id': row['origin_id'],
                 'huohao': row['huohao'], 'status': 'exception', 'error': str(e)[:200]}
        results.append(r)
        print(f"    -> {r['status']}" + (f" | {r.get('error', '')[:80]}" if r.get('error') else ""))
        for s in r.get('steps', []):
            flag = 'OK ' if s['ok'] else 'XX '
            print(f"      {flag}{s['name']}" + (f" | {str(s.get('extra', ''))[:80]}" if s.get('extra') else ""))
        # 每行处理完关闭 publish/success 页（只保留在售页），继续下一行
        from clone_id_flow import close_processed_tabs
        close_processed_tabs(args.cdp_url)

    # 报告先落盘（write_back 可能因文件占用/权限失败，不能丢新ID）
    if args.report:
        Path(args.report).write_text(json.dumps({'summary': {
            'total': len(results),
            'success': sum(1 for r in results if r['status'] == 'success'),
            'failed': sum(1 for r in results if r['status'] in ('failed', 'exception')),
            'unknown': sum(1 for r in results if r['status'] == 'unknown_error_recorded')},
            'results': results}, ensure_ascii=False, indent=1), encoding='utf-8')
        print('[report]', args.report)

    # 写回 D=新建id, E=状态
    if args.yes:
        changed = 0
        for r in results:
            idx = r['excel_row']
            st = r['status']
            if st == 'success':
                ws.cell(idx, 4).value = r.get('new_id', '')
                ws.cell(idx, 5).value = f"已发布 {datetime.now().strftime('%m-%d %H:%M')}"
            elif st == 'failed':
                ws.cell(idx, 5).value = f"失败: {r.get('error', '')[:40]}"
            elif st == 'unknown_error_recorded':
                ws.cell(idx, 5).value = f"未知错误: {r.get('error', '')[:30]}"
            changed += 1
        wb.save(args.src)
        print(f'\n[write-back] {changed} 行写回 {args.src}')
    else:
        print('\n[write-back] 跳过（未 --yes）')


if __name__ == '__main__':
    main()
