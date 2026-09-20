# -*- coding: utf-8 -*-
"""校验：①生成的留档 xls vs API JSON；②新主表新增块 vs 留档

用法:
    py verify_all.py --base <基准主表.xlsx> --master <新主表.xlsx>
                     [--archive-dir <留档目录>] [--json-dir <api_json目录>] [--days ...]

    新增块起始行 = 基准主表末行 + 1（从基准 zip 的 dimension 自动推导）。
"""
import xlrd, json, os, re, sys, argparse, zipfile
from openpyxl import load_workbook

DEFAULT_ARCHIVE_DIR = os.environ.get("SYCM_SEARCH_ARCHIVE_DIR", r"E:\巴拉巴拉\搜索词数据留存")


def detect_last_row(xlsx_path):
    with zipfile.ZipFile(xlsx_path, "r") as z:
        sheet_xml = z.read("xl/worksheets/sheet1.xml").decode("utf-8")
    m = re.search(r'<dimension ref="A1:L(\d+)"', sheet_xml)
    if not m:
        raise RuntimeError(f"未能在 {xlsx_path} 中找到 dimension")
    return int(m.group(1))


def parse_archive_xls(archive_dir, date_str):
    fn = os.path.join(archive_dir, f"【生意参谋】选词助手-引流搜索词-店外-无线-{date_str.replace('-', '')}.xls")
    book = xlrd.open_workbook(fn)
    sh = book.sheet_by_index(0)
    return [[sh.cell_value(r, c) for c in range(10)] for r in range(6, sh.nrows)]


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--base", required=True, help="基准主表 xlsx（未拼接前）")
    ap.add_argument("--master", required=True, help="新主表 xlsx（拼接后）")
    ap.add_argument("--archive-dir", default=DEFAULT_ARCHIVE_DIR)
    ap.add_argument("--json-dir", default=os.path.join(DEFAULT_ARCHIVE_DIR, "api_json"))
    ap.add_argument("--days", default="", help="逗号分隔日期；缺省自动扫描 json-dir")
    args = ap.parse_args()

    if args.days:
        days = [d.strip() for d in args.days.split(",") if d.strip()]
    else:
        days = sorted(
            f[4:8] + "-" + f[8:10] + "-" + f[10:12]
            for f in os.listdir(args.json_dir)
            if re.fullmatch(r"api_\d{8}\.json", f)
        )

    # ---------- 1. 留档 xls vs API JSON ----------
    for d in days:
        fn = os.path.join(args.archive_dir, f"【生意参谋】选词助手-引流搜索词-店外-无线-{d.replace('-', '')}.xls")
        if not os.path.exists(fn):
            print(d, "留档不存在，跳过 API 比对")
            continue
        book = xlrd.open_workbook(fn)
        sh = book.sheet_by_index(0)
        xls_rows = [[sh.cell_value(r, c) for c in range(10)] for r in range(6, sh.nrows)]
        jf = os.path.join(args.json_dir, f"api_{d.replace('-', '')}.json")
        if not os.path.exists(jf):
            print(d, "留档行数", len(xls_rows), "(无 API JSON，跳过比对)")
            continue
        res = json.load(open(jf, encoding="utf-8"))
        api_rows = res["rows"]
        print(d, "留档", len(xls_rows), "vs API", len(api_rows), "首词:", xls_rows[0][1], "|", api_rows[0]["seKeyword"]["value"])
        ok = True
        for i in [0, 1, 2, len(xls_rows) - 1]:
            x = xls_rows[i]
            a = api_rows[i]
            if x[1] != a.get("seKeyword", {}).get("value"):
                ok = False
                print("  WORD DIFF", i, x[1], a.get("seKeyword", {}).get("value"))
        if not ok:
            print("  ^^ 存在差异")

    # ---------- 2. 新主表新增块 vs 留档 ----------
    base_last = detect_last_row(args.base)
    start_row = base_last + 1
    wb = load_workbook(args.master, read_only=True, data_only=True)
    ws = wb["Sheet1"]
    new_block = []
    for row in ws.iter_rows(min_row=start_row, values_only=True):
        new_block.append(["" if v is None else str(v) for v in row])
    wb.close()
    print("新表新增块行数:", len(new_block), "起始行:", start_row)

    # 主表列映射：col0年份 col1月份 col2日期 col3词 col4访客 col5加购 col6收藏 col7买家 col8率 col9金额 col10客单价 col11UV
    idx = 0
    mismatches = 0
    for d in days:
        if not os.path.exists(os.path.join(args.archive_dir, f"【生意参谋】选词助手-引流搜索词-店外-无线-{d.replace('-', '')}.xls")):
            continue
        arch = parse_archive_xls(args.archive_dir, d)
        for i, a10 in enumerate(arch):
            m = new_block[idx]
            # 留档行 a10 = [日期, 词, 指标...]，主表行 = [年份, 月份] + a10
            expect = [f"{d[:4]}年", f"{int(d[5:7])}月"] + [str(v) for v in a10]
            if m != expect:
                mismatches += 1
                if mismatches <= 5:
                    print("MISMATCH", d, i, m, "VS", expect)
            idx += 1
    print("主表新增块 vs 留档: 不一致数 =", mismatches, "/", len(new_block))


if __name__ == "__main__":
    main()
